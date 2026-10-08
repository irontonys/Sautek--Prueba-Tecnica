"""Lectura de la cotización de un proveedor por códigos de producto, sin IA.

Se saca el texto del archivo (PDF, Excel o CSV) y se toma como partida cada renglón
que trae un código de producto y al menos un monto. Luego se compara contra la RFQ.
"""

import csv
import re
import subprocess
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from resurtido.orders import money
from resurtido.validation import normalize_code

IGUAL = "Igual"
CAMBIO_PRECIO = "Cambio de precio"
PARCIAL = "Surtido parcial"
CAMBIO_CANTIDAD = "Cambio de cantidad"
SIN_EXISTENCIA = "Sin existencia"
NO_COTIZADO = "No viene en la cotización"
NO_EN_RFQ = "No está en la RFQ"

CODE = re.compile(r"\bFTR-?\d{1,4}\b", re.IGNORECASE)
# Montos: con signo de pesos o con decimales ("$1,250.00", "559.58", "33574.8").
NUMBER = re.compile(r"\$?\s?\d[\d,]*(?:\.\d+)?")
OUT_OF_STOCK = re.compile(r"sin\s+existencia|agotado|no\s+disponible|sin\s+stock", re.IGNORECASE)
TOLERANCE = Decimal("0.01")


class QuoteReadError(Exception):
    """El archivo no se pudo leer como cotización."""


@dataclass
class QuotedLine:
    codigo: str
    cantidad: Decimal
    precio: Decimal


@dataclass
class ComparisonRow:
    codigo: str
    descripcion: str
    cantidad_rfq: Decimal
    precio_rfq: Decimal
    cantidad_cotizada: Decimal
    precio_cotizado: Decimal
    estatus: str


def extract_text(path):
    """Texto del archivo, un renglón de la cotización por línea."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            result = subprocess.run(
                ["pdftotext", "-layout", str(path), "-"], capture_output=True, check=True, timeout=60
            )
        except FileNotFoundError as exc:
            raise QuoteReadError("No está instalado pdftotext (poppler-utils) para leer PDF") from exc
        except subprocess.CalledProcessError as exc:
            raise QuoteReadError(f"No se pudo leer el PDF: {exc.stderr.decode(errors='replace').strip()}") from exc
        return result.stdout.decode("utf-8", errors="replace")
    if suffix == ".xlsx":
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise QuoteReadError(f"No se pudo leer el Excel: {exc}") from exc
        lines = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                lines.append("\t".join(_cell_text(value) for value in row if value is not None))
        workbook.close()
        return "\n".join(lines)
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", errors="replace", newline="") as handle:
            return "\n".join("\t".join(row) for row in csv.reader(handle))
    raise QuoteReadError(f"No se pueden leer archivos {suffix or 'sin extensión'}: solo PDF, Excel (.xlsx) o CSV")


def _cell_text(value):
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:.2f}"
    return str(value)


def _to_decimal(token):
    try:
        return Decimal(token.replace("$", "").replace(",", "").strip())
    except InvalidOperation:
        return None


def _is_amount(token):
    return "$" in token or "." in token


def parse_line(line):
    """Una partida si el renglón trae un código y al menos un monto; si no, None."""
    match = CODE.search(line)
    if not match:
        return None
    rest = line[match.end():]
    tokens = [t.strip() for t in NUMBER.findall(rest)]
    amounts = [i for i, t in enumerate(tokens) if _is_amount(t)]
    if not amounts:
        return None  # un párrafo que menciona códigos, no una partida
    price = _to_decimal(tokens[amounts[0]])
    total = _to_decimal(tokens[amounts[-1]]) if len(amounts) > 1 else None
    quantities = [_to_decimal(t) for t in tokens[:amounts[0]] if not _is_amount(t)]
    quantities = [q for q in quantities if q is not None]
    if OUT_OF_STOCK.search(rest):
        quantity = Decimal("0")
    elif total is not None and any(abs(q * price - total) <= TOLERANCE for q in quantities):
        # La cantidad surtible es la que cuadra con el importe (no la solicitada ni la de la descripción).
        quantity = next(q for q in reversed(quantities) if abs(q * price - total) <= TOLERANCE)
    elif quantities:
        quantity = quantities[-1]
    else:
        return None
    code, _ = normalize_code(match.group(0))
    return QuotedLine(code, quantity, money(price))


def parse_quote(text):
    """Regresa {código: QuotedLine}; si un código aparece dos veces se queda la primera partida."""
    quoted = {}
    for line in text.splitlines():
        item = parse_line(line)
        if item and item.codigo not in quoted:
            quoted[item.codigo] = item
    return quoted


def compare(rfq_lines, quoted):
    """`rfq_lines`: [(código, descripción, cantidad, precio)]. Un renglón por producto de ambos lados."""
    rows = []
    for codigo, descripcion, cantidad, precio in rfq_lines:
        cantidad, precio = Decimal(str(cantidad)), money(precio)
        item = quoted.get(codigo)
        if item is None:
            rows.append(ComparisonRow(codigo, descripcion, cantidad, precio, Decimal("0"), Decimal("0"), NO_COTIZADO))
            continue
        if item.cantidad == 0:
            status = SIN_EXISTENCIA
        elif item.cantidad < cantidad:
            status = PARCIAL
        elif item.cantidad > cantidad:
            status = CAMBIO_CANTIDAD
        elif abs(item.precio - precio) > TOLERANCE:
            status = CAMBIO_PRECIO
        else:
            status = IGUAL
        rows.append(ComparisonRow(codigo, descripcion, cantidad, precio, item.cantidad, item.precio, status))
    in_rfq = {line[0] for line in rfq_lines}
    rows += [
        ComparisonRow(item.codigo, "", Decimal("0"), Decimal("0"), item.cantidad, item.precio, NO_EN_RFQ)
        for codigo, item in quoted.items() if codigo not in in_rfq
    ]
    return rows
