"""Punto de entrada por línea de comandos: `python -m resurtido`."""

import argparse
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from resurtido.loader import ColumnError, load_sheets
from resurtido.orders import build_orders, order_summary_lines
from resurtido.report import summary_lines, write_exceptions, write_orders
from resurtido.validation import sort_exceptions, validate

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_INPUT = REPO_ROOT / "data" / "inventario_ferreteria_garza.xlsx"
DEFAULT_OUTPUT = REPO_ROOT / "output"

EXPECTED_SHEETS = ["Existencias", "Minimos", "Producto_Proveedor", "Proveedores"]


class InputError(Exception):
    """Problema con el archivo de entrada que se reporta sin traceback."""


def parse_args(argv):
    parser = argparse.ArgumentParser(
        prog="python -m resurtido",
        description="Arma los pedidos de resurtido de la semana a partir del Excel de inventario.",
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=DEFAULT_INPUT,
        help="Excel de entrada (por omisión: data/inventario_ferreteria_garza.xlsx)",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help="Carpeta de salida (por omisión: output/)",
    )
    return parser.parse_args(argv)


def read_sheet_names(path):
    if not path.exists():
        raise InputError(f"No existe el archivo de entrada: {path}")
    if not path.is_file():
        raise InputError(f"La ruta de entrada no es un archivo: {path}")
    try:
        workbook = load_workbook(path, read_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise InputError(f"No se pudo leer como Excel (.xlsx): {path} ({exc})") from exc
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


def check_sheets(sheet_names):
    missing = [name for name in EXPECTED_SHEETS if name not in sheet_names]
    if missing:
        raise InputError("Faltan hojas en el Excel: " + ", ".join(missing))


def main(argv=None):
    args = parse_args(argv)
    try:
        check_sheets(read_sheet_names(args.input))
        sheets = load_sheets(args.input)
    except (InputError, ColumnError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1

    result = validate(sheets)
    orders, order_exceptions = build_orders(result.products, result.suppliers)
    result.exceptions = sort_exceptions(result.exceptions + order_exceptions)
    args.output.mkdir(parents=True, exist_ok=True)
    report = write_exceptions(result.exceptions, args.output)
    orders_file = write_orders(orders, args.output)

    print(f"Excel leído: {args.input}")
    print("Hojas encontradas: " + ", ".join(EXPECTED_SHEETS))
    for line in summary_lines(result) + order_summary_lines(orders):
        print(line)
    print(f"Pedidos: {orders_file}")
    print(f"Reporte de excepciones: {report}")
    return 0
