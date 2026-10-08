"""Lectura de las cuatro hojas del Excel de inventario, sin interpretar los valores."""

import pandas as pd
from openpyxl import load_workbook
from pandas._libs.parsers import STR_NA_VALUES

SHEET_COLUMNS = {
    "Existencias": ["Codigo", "Descripcion", "Existencia", "Unidad", "Estatus"],
    "Minimos": ["Codigo", "Minimo", "Maximo"],
    "Producto_Proveedor": ["Codigo", "Proveedor_ID", "Costo_unitario_MXN", "Empaque_multiplo"],
    "Proveedores": ["Proveedor_ID", "Nombre", "Correo", "Pedido_minimo_MXN"],
}

# Fila 1: título descriptivo. Fila 2: encabezados. Fila 3 en adelante: datos.
HEADER_ROW_INDEX = 1
FIRST_DATA_ROW = 3

ROW_COLUMN = "fila_excel"


class ColumnError(Exception):
    """Falta una columna esperada en alguna hoja."""


def load_sheets(path):
    """Regresa {hoja: DataFrame} con valores crudos y la columna `fila_excel`.

    Los valores se leen como objeto (sin inferir tipos) para poder reportar
    exactamente lo que trae cada celda, por ejemplo `"1,250"` guardado como texto.
    """
    raw = read_raw_sheets(path, list(SHEET_COLUMNS))
    sheets = {}
    for name, columns in SHEET_COLUMNS.items():
        frame = raw[name]
        frame.columns = [str(column).strip() for column in frame.columns]
        missing = [column for column in columns if column not in frame.columns]
        if missing:
            raise ColumnError(f"En la hoja {name} falta la columna: " + ", ".join(missing))
        frame = frame[columns].copy()
        frame[ROW_COLUMN] = frame.index + FIRST_DATA_ROW
        # Un renglón totalmente vacío no es un registro: es espacio en blanco del Excel.
        frame = frame.dropna(how="all", subset=columns).reset_index(drop=True)
        sheets[name] = frame
    return sheets


def read_raw_sheets(path, names):
    """Lee las celdas con openpyxl, cada una con su tipo.

    `pd.read_excel` infiere tipos aunque se le pida `dtype=object`: en una columna con
    `True` y `1` convierte el `1` en `True`. Se conservan sus otras dos conversiones, de
    las que dependen las reglas: un flotante entero (`116.0`) se lee como `int` y los
    textos que pandas toma como vacío (`""`, `"NA"`, `"#N/A"`...) se leen como vacío.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        frames = {}
        for name in names:
            rows = [list(row) for row in workbook[name].iter_rows(values_only=True)]
            header = rows[HEADER_ROW_INDEX] if len(rows) > HEADER_ROW_INDEX else []
            width = len(header)
            columns = [
                f"Unnamed: {index}" if column is None else column
                for index, column in enumerate(header)
            ]
            data = [
                [_cell_value(value) for value in (row + [None] * width)[:width]]
                for row in rows[HEADER_ROW_INDEX + 1:]
            ]
            frames[name] = pd.DataFrame(data, columns=columns, dtype=object)
        return frames
    finally:
        workbook.close()


def _cell_value(value):
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and value in STR_NA_VALUES:
        return None
    return value
