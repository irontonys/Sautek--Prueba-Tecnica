"""Escritura del reporte de excepciones."""

import csv
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Font

EXCEPTIONS_FILE = "excepciones.csv"
ORDERS_FILE = "pedidos.xlsx"
MONEY_FORMAT = '"$"#,##0.00'

SUMMARY_COLUMNS = [
    ("Proveedor_ID", 13), ("Proveedor", 32), ("Correo", 28), ("Productos", 11),
    ("Total_MXN", 15), ("Pedido_minimo_MXN", 19), ("Estado", 24),
]
DETAIL_COLUMNS = [
    ("Proveedor_ID", 13), ("Proveedor", 32), ("Codigo", 11), ("Descripcion", 30),
    ("Existencia", 11), ("Minimo", 9), ("Maximo", 9), ("Empaque_multiplo", 17),
    ("Cantidad", 10), ("Costo_unitario_MXN", 19), ("Importe_MXN", 15),
    ("Estado_pedido", 24), ("Revisar", 9),
]
EXCEPTION_COLUMNS = ["hoja", "fila_excel", "codigo", "tipo", "detalle", "valor_original", "decision"]


def write_exceptions(exceptions, output_dir):
    path = output_dir / EXCEPTIONS_FILE
    # utf-8-sig agrega el BOM para que Excel en Windows muestre bien los acentos.
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(EXCEPTION_COLUMNS)
        for item in exceptions:
            writer.writerow([getattr(item, column) for column in EXCEPTION_COLUMNS])
    return path


def summary_lines(result):
    lines = [
        f"Renglones leídos en Existencias: {result.rows_read}",
        f"  Repetidos (se conserva el primero): {result.repeated}",
        f"  Procesables: {len(result.products)}",
        f"  No procesables: {len(result.excluded)}",
        f"Excepciones: {len(result.exceptions)}",
    ]
    for tipo, count in sorted(Counter(e.tipo for e in result.exceptions).items()):
        lines.append(f"  {tipo}: {count}")
    return lines


def _add_sheet(workbook, title, columns, rows, money_columns):
    sheet = workbook.create_sheet(title)
    sheet.append([name for name, _ in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, (_, width) in enumerate(columns, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    names = [name for name, _ in columns]
    for row in rows:
        sheet.append(row)
        for name in money_columns:
            sheet.cell(sheet.max_row, names.index(name) + 1).number_format = MONEY_FORMAT
    sheet.freeze_panes = "A2"


def write_orders(orders, output_dir):
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = [
        [
            o.supplier.proveedor_id, o.supplier.nombre, o.supplier.correo, len(o.lines),
            float(o.total), o.supplier.pedido_minimo, o.estado,
        ]
        for o in orders
    ]
    detail = [
        [
            o.supplier.proveedor_id, o.supplier.nombre, line.product.codigo,
            line.product.descripcion, line.product.existencia, line.product.minimo,
            line.product.maximo, line.product.empaque_multiplo, line.cantidad,
            line.product.costo_unitario, float(line.importe), o.estado,
            "Sí" if line.product.revisar else "",
        ]
        for o in orders
        for line in o.lines
    ]
    _add_sheet(workbook, "Resumen", SUMMARY_COLUMNS, summary, ["Total_MXN", "Pedido_minimo_MXN"])
    _add_sheet(workbook, "Detalle", DETAIL_COLUMNS, detail, ["Costo_unitario_MXN", "Importe_MXN"])
    path = output_dir / ORDERS_FILE
    workbook.save(path)
    return path
