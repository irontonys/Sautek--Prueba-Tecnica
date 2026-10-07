"""Escenarios de la capacidad `calculo-pedidos`."""

from decimal import Decimal

import pytest
from openpyxl import load_workbook

from resurtido import cli
from resurtido.loader import load_sheets
from resurtido.orders import (
    NO_SE_ENVIA, SE_ENVIA, SIN_PRODUCTOS, build_orders, order_quantity,
)
from resurtido.report import write_orders
from resurtido.validation import PEDIDO_MINIMO_NO_ALCANZADO, Product, Supplier, validate


def make_product(codigo="FTR-0001", existencia=3, minimo=30, maximo=60, multiplo=6,
                 costo=100.0, proveedor="P01", revisar=False):
    return Product(
        codigo=codigo, descripcion=f"Producto {codigo}", existencia=existencia, minimo=minimo,
        maximo=maximo, proveedor_id=proveedor, costo_unitario=costo, empaque_multiplo=multiplo,
        fila_excel=3, revisar=revisar,
    )


def make_suppliers(pedido_minimo=0, **extra):
    suppliers = {"P01": Supplier("P01", "Proveedor Uno", "uno@prueba.mx", pedido_minimo, 3)}
    for supplier_id, minimo in extra.items():
        suppliers[supplier_id] = Supplier(supplier_id, f"Proveedor {supplier_id}", "x@prueba.mx", minimo, 4)
    return suppliers


def single_order(products, pedido_minimo=0):
    orders, exceptions = build_orders(products, make_suppliers(pedido_minimo))
    return orders[0], exceptions


# Selección


def test_existencia_bajo_el_minimo_se_pide():
    order, _ = single_order([make_product(existencia=3, minimo=30)])

    assert [line.product.codigo for line in order.lines] == ["FTR-0001"]


def test_existencia_igual_al_minimo_no_se_pide():
    order, _ = single_order([make_product(existencia=30, minimo=30)])

    assert order.lines == []
    assert order.estado == SIN_PRODUCTOS


def test_producto_no_procesable_no_aparece():
    sheets = load_sheets(cli.DEFAULT_INPUT)
    result = validate(sheets)
    orders, _ = build_orders(result.products, result.suppliers)

    ordered = {line.product.codigo for o in orders for line in o.lines}
    assert ordered.isdisjoint(result.excluded)
    assert "FTR-0019" not in ordered  # mínimo mayor que máximo, aunque esté bajo el mínimo


# Cantidad


@pytest.mark.parametrize(
    "existencia, maximo, multiplo, esperado",
    [
        (3, 60, 6, 60),     # faltan 57 → siguiente múltiplo de 6
        (0, 20, 1, 20),     # exacto
        (11, 80, 12, 72),   # faltan 69 → 72
        (2.5, 10, 1, 8),    # faltan 7.5 → se sube a 8
        (0, 24, 12, 24),    # ya es múltiplo
    ],
)
def test_cantidad_redondeada_al_multiplo(existencia, maximo, multiplo, esperado):
    assert order_quantity(existencia, maximo, multiplo) == esperado


def test_existencia_negativa_ajustada_pide_hasta_el_maximo():
    order, _ = single_order([make_product(existencia=0, minimo=12, maximo=24, multiplo=1, revisar=True)])

    assert order.lines[0].cantidad == 24


# Agrupación y total


def test_total_del_proveedor_al_centavo():
    products = [
        make_product("FTR-0001", existencia=0, maximo=1, multiplo=1, costo=1000.50),
        make_product("FTR-0002", existencia=0, maximo=1, multiplo=1, costo=2000.25),
    ]

    order, _ = single_order(products)

    assert order.total == Decimal("3000.75")
    assert order.total == sum(line.importe for line in order.lines)


def test_productos_agrupados_por_proveedor():
    products = [
        make_product("FTR-0001", proveedor="P01"),
        make_product("FTR-0002", proveedor="P02"),
        make_product("FTR-0003", proveedor="P01"),
    ]

    orders, _ = build_orders(products, make_suppliers(P02=0))

    by_id = {o.supplier.proveedor_id: [l.product.codigo for l in o.lines] for o in orders}
    assert by_id == {"P01": ["FTR-0001", "FTR-0003"], "P02": ["FTR-0002"]}


# Pedido mínimo


def test_no_alcanza_el_minimo():
    product = make_product(existencia=0, maximo=1, multiplo=1, costo=12664.05)

    order, exceptions = single_order([product], pedido_minimo=15000)

    assert order.estado == NO_SE_ENVIA
    [found] = exceptions
    assert found.tipo == PEDIDO_MINIMO_NO_ALCANZADO
    assert "2,335.95" in found.detalle
    assert "No se envía" in found.decision


def test_alcanza_el_minimo():
    product = make_product(existencia=0, maximo=1, multiplo=1, costo=5000)

    order, exceptions = single_order([product], pedido_minimo=5000)

    assert order.estado == SE_ENVIA
    assert exceptions == []


# Archivo de pedidos


def test_archivo_de_pedidos(tmp_path):
    products = [make_product(revisar=True), make_product("FTR-0002", existencia=50)]
    orders, _ = build_orders(products, make_suppliers(P02=0))

    workbook = load_workbook(write_orders(orders, tmp_path))

    assert workbook.sheetnames == ["Resumen", "Detalle"]
    resumen = list(workbook["Resumen"].values)
    assert resumen[0][-1] == "Estado"
    assert [row[-1] for row in resumen[1:]] == [SE_ENVIA, SIN_PRODUCTOS]
    detalle = list(workbook["Detalle"].values)
    assert len(detalle) == 2
    assert detalle[1][2] == "FTR-0001" and detalle[1][-1] == "Sí"


# Regresión sobre el Excel real


@pytest.fixture(scope="module")
def real_orders():
    result = validate(load_sheets(cli.DEFAULT_INPUT))
    orders, exceptions = build_orders(result.products, result.suppliers)
    return {o.supplier.proveedor_id: o for o in orders}, exceptions


def test_totales_por_proveedor_excel_real(real_orders):
    orders, _ = real_orders

    assert {sid: (str(o.total), o.estado) for sid, o in orders.items()} == {
        "P01": ("168209.93", SE_ENVIA),
        "P02": ("22912.98", SE_ENVIA),
        "P03": ("44175.54", SE_ENVIA),
        "P04": ("10201.20", SE_ENVIA),
        "P05": ("12664.05", NO_SE_ENVIA),
        "P06": ("26098.32", SE_ENVIA),
        "P07": ("39531.63", SE_ENVIA),
        "P08": ("0.00", SIN_PRODUCTOS),
    }


def test_segueta_marcada_para_revision(real_orders):
    orders, _ = real_orders

    [line] = [l for o in orders.values() for l in o.lines if l.product.codigo == "FTR-0007"]
    assert line.product.revisar
    assert line.product.existencia == 0


def test_corrida_completa_genera_pedidos(tmp_path, capsys):
    assert cli.main(["--output", str(tmp_path)]) == 0

    out = capsys.readouterr().out
    assert "Pedidos que se envían: 6" in out
    assert "Pedidos que no se envían (no alcanzan el pedido mínimo): 1" in out
    resumen = list(load_workbook(tmp_path / "pedidos.xlsx")["Resumen"].values)
    assert len(resumen) == 9  # encabezado + 8 proveedores
    assert "Pedido mínimo no alcanzado" in (tmp_path / "excepciones.csv").read_text(encoding="utf-8-sig")
